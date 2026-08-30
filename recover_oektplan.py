from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


OUT_DIR = Path(r"C:\Users\Madsf\OneDrive\Belgeler")
TXT_PATH = OUT_DIR / "Gjenopprettet_Oekt_Plan_for_barn.txt"
DOCX_PATH = OUT_DIR / "Gjenopprettet_Oekt_Plan_for_barn.docx"

CONTENT = """ØKTPLAN FOR BARN PÅ 5 ÅR - 45 MINUTTER

Tema
Lek, bevegelsesglede og mestring

Målgruppe
Barn på 5 år

Varighet
45 minutter

Utstyr
Kjegler, erteposer eller baller, matter, benker, ringer, tau eller annet enkelt hinderløypeutstyr.

Mål for økten
- Barna skal oppleve bevegelsesglede gjennom lek tilpasset 5-åringer.
- Barna skal øve på å løpe, stoppe, balansere, krype, hoppe og samarbeide.
- Alle skal kunne delta på sitt nivå og kjenne mestring.

Kort gjennomføring
Samle barna først og forklar kort at økten består av tre leker. De tre første lekene varer omtrent 10 minutter hver. De siste 15 minuttene brukes til hinderløype.

1. Lek: Haien kommer - 10 minutter
Organisering: Marker et område som er havet. Legg gjerne ut matter, ringer eller kjegler som øyer. Ett barn eller en voksen er hai.
Slik gjør dere: Barna beveger seg rundt i havet. Når den voksne roper «Haien kommer!», skal barna prøve å komme seg til en øy før haien tar dem. Blir et barn tatt, kan barnet hjelpe haien, eller prøve på nytt fra siden etter noen sekunder.
Tilpasning: For 5-åringer kan en voksen gjerne være hai først. Hold tempoet lekent og trygt, og la barna få god tid til å finne en øy.
Fokus: Løping, reaksjon, romforståelse og lekeglede.

2. Lek: Fargeleken - 10 minutter
Organisering: Legg kjegler, ringer eller lapper i ulike farger rundt i rommet. Barna står samlet i midten.
Slik gjør dere: Den voksne roper en farge, for eksempel «rød!». Barna skal da løpe, gå, hoppe eller krype til riktig farge.
Varier bevegelsene: Gå som en kjempe, liste som en mus, hoppe som en frosk, krype som en soldat eller løpe rolig til fargen.
Tilpasning: Bruk 4-6 tydelige farger. Hjelp barna ved å peke eller vise fargen først hvis noen blir usikre.
Fokus: Bevegelse, fargegjenkjenning, lytting og fantasi.

3. Stafett - 10 minutter
Organisering: Del barna inn i 3 grupper.
Slik gjør dere: Første barn i hver gruppe kaster en gjenstand oppi en ring. Hvis barnet treffer, løper det fra den ene siden til den andre og tilbake igjen. Når barnet er tilbake, er det nestemann sin tur. Leken fortsetter til alle i gruppen har vært gjennom.

4. Reserveleker
1. Rødt lys
2. Tikken
3. Stiv heks
4. Bjørnen sover

5. Hinderløype - 15 minutter
Organisering: Lag en enkel hinderløype med 4-6 stasjoner. Barna går gjennom løypa én og én eller i små grupper.
Balanse: Barna skal gå over og balansere på en benk.
Klatre: Barna skal klatre på ribbevegg.
Kjegler: Barna skal løpe sikksakk rundt kjegler.
Rokkering: Barna skal hoppe fra den ene rokkeringen til den andre.
Krype: Barna skal krype under noe.
Over: Barna skal klatre over noe.

Avslutning - 2 minutter hvis det passer
Samle barna i en sirkel.
Spør gjerne: Hva var morsomst? Hva klarte du i hinderløypa?
Avslutt med ros for innsats og deltakelse.
"""


def add_paragraph(doc: Document, text: str) -> None:
    if not text:
        return
    if text.isupper() and len(text) > 10:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(20)
        p.paragraph_format.space_after = Pt(8)
    elif text in {"Tema", "Målgruppe", "Varighet", "Utstyr", "Mål for økten", "Kort gjennomføring"} or text[:2].isdigit() or text.startswith(("1. Lek", "2. Lek", "3. ", "4. ", "5. ", "Avslutning")):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    elif text.startswith("- "):
        p = doc.add_paragraph(text[2:], style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
    else:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(4)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TXT_PATH.write_text(CONTENT, encoding="utf-8")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    for line in CONTENT.splitlines():
        add_paragraph(doc, line.strip())
    doc.save(DOCX_PATH)
    print(TXT_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
